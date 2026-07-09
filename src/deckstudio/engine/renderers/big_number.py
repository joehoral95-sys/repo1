"""Big-number slide: 1-4 KPI stat tiles. One stat = hero treatment."""

from __future__ import annotations

from ...spec.schema import BigNumberSlide
from ..geometry import Box, columns, vcenter
from ..infographics import add_stat_tile
from ..registry import renderer
from ._common import add_title_band


@renderer("big_number")
def render(slide, model: BigNumberSlide, ctx) -> None:
    """Variants: cards (navy hero first), uniform (all light), dark (all navy)."""
    tokens = ctx.tokens
    variant, chrome = ctx.variant_parts(model)
    area = add_title_band(slide, tokens, model.title, kicker=model.kicker,
                          style=chrome)
    n = len(model.stats)
    if n == 1:
        _solo_hero(slide, area, model.stats[0], tokens)
        return
    tile_h = min(3.2, area.height_in - 0.3)
    for i, (stat, col) in enumerate(
            zip(model.stats, columns(area, n, tokens.gutter_in), strict=True)):
        hero = (variant == "dark") or (variant == "cards" and i == 0)
        add_stat_tile(slide, vcenter(col, tile_h), stat, tokens,
                      hero=hero, accent_index=i - 1)


def _solo_hero(slide, area, stat, tokens) -> None:
    """One stat = the whole stage: giant blue number, swash, delta chip."""
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Pt

    from ..infographics import ARROWS
    from ..shapes import add_brand_art, add_chip, chip_width
    from ..text import add_text

    cx = area.left_in
    value_box = Box(cx, area.top_in + 0.25, area.width_in, 1.9)
    shape = add_text(slide, value_box, stat.value, tokens, scale="big_number",
                     role="heading", color="accent", bold=True,
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                     shrink_to_fit=True)
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(120)
    add_brand_art(slide, tokens, "swash_sky",
                  Box(cx + area.width_in / 2 - 1.5, value_box.bottom_in + 0.05, 3.0, 0.38))
    add_text(slide, Box(cx, value_box.bottom_in + 0.55, area.width_in, 0.5),
             stat.label, tokens, scale="subtitle", color="neutral_dark",
             align=PP_ALIGN.CENTER)
    if stat.delta:
        text = f"{ARROWS[stat.arrow]} {stat.delta}".strip()
        chip_w = chip_width(text)
        fill = {"good": "positive", "bad": "negative", "neutral": "neutral_mid"}[stat.sentiment]
        add_chip(slide, cx + (area.width_in - chip_w) / 2, value_box.bottom_in + 1.15,
                 text, tokens, fill=fill, text_color="white")
